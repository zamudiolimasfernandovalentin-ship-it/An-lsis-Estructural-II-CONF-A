from __future__ import annotations

import io
from typing import Any, Dict, Iterable

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Flowable, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _records_table(sheet, title: str, columns: list[str], records: Iterable[dict]) -> None:
    sheet.append([title])
    sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=max(1, len(columns)))
    sheet.append(columns)
    for record in records:
        sheet.append([record.get(column, "") for column in columns])


def _format_workbook(workbook: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    title_fill = PatternFill("solid", fgColor="24445C")
    border = Border(bottom=Side(style="thin", color="B8C5CF"))
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A3" if sheet.max_row > 1 else "A1"
        sheet.sheet_view.showGridLines = False
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(cell.value, (int, float)) else "left")
                cell.border = border
        if sheet.max_row:
            if sheet.title in {"K_Global", "Particion"}:
                for row in sheet.iter_rows():
                    for cell in row:
                        cell.fill = PatternFill(fill_type=None)
                        cell.font = Font(bold=cell.row <= 2, color="000000")
            else:
                for cell in sheet[1]:
                    cell.fill = title_fill
                    cell.font = Font(bold=True, color="FFFFFF")
                if sheet.max_row > 1:
                    for cell in sheet[2]:
                        cell.fill = header_fill
                        cell.font = Font(bold=True, color="183247")
        for column_index in range(1, sheet.max_column + 1):
            column_cells = next(sheet.iter_cols(min_col=column_index, max_col=column_index))
            column_letter = get_column_letter(column_index)
            width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 28)
            sheet.column_dimensions[column_letter].width = width


def _matrix_table(matrix: Any, row_labels: list[str], column_labels: list[str], available_width: float) -> LongTable:
    _, column_count = np.asarray(matrix).shape
    first_col_width = min(max(available_width * 0.1, 38), 62)
    cell_width = (available_width - first_col_width) / max(column_count, 1)
    font_size = min(7.0, max(4.0, cell_width * 0.18))
    data = [["GDL", *column_labels]]
    data.extend([[row_labels[index], *[f"{float(value):.4g}" for value in row]] for index, row in enumerate(matrix)])
    table = LongTable(data, colWidths=[first_col_width] + [cell_width] * column_count, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 1), (-1, -1), "Courier"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#aab8c2")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


class MemberDiagram(Flowable):
    def __init__(self, element_id: int, node_i: int, node_j: int, x_i: float, y_i: float, x_j: float, y_j: float, phi: float, width: float):
        super().__init__()
        self.element_id = element_id
        self.node_i = node_i
        self.node_j = node_j
        self.x_i = x_i
        self.y_i = y_i
        self.x_j = x_j
        self.y_j = y_j
        self.phi = phi
        self.width = width
        self.height = 132

    def _arrow(self, start_x: float, start_y: float, end_x: float, end_y: float, label: str) -> None:
        drawing = self.canv
        dx, dy = end_x - start_x, end_y - start_y
        length = max(float(np.hypot(dx, dy)), 1e-9)
        ux, uy = dx / length, dy / length
        px, py = -uy, ux
        drawing.line(start_x, start_y, end_x, end_y)
        drawing.line(end_x, end_y, end_x - ux * 5 + px * 2.5, end_y - uy * 5 + py * 2.5)
        drawing.line(end_x, end_y, end_x - ux * 5 - px * 2.5, end_y - uy * 5 - py * 2.5)
        drawing.setFont("Helvetica", 7)
        drawing.drawString(end_x + 3, end_y + 2, label)

    def draw(self) -> None:
        drawing = self.canv
        dx, dy = self.x_j - self.x_i, self.y_j - self.y_i
        length = float(np.hypot(dx, dy))
        horizontal_scale = (self.width - 170) / abs(dx) if abs(dx) > 1e-9 else float("inf")
        vertical_scale = (self.height - 54) / abs(dy) if abs(dy) > 1e-9 else float("inf")
        scale = min(horizontal_scale, vertical_scale)
        center_x, center_y = self.width / 2, self.height / 2
        start_x, start_y = center_x - dx * scale / 2, center_y - dy * scale / 2
        end_x, end_y = center_x + dx * scale / 2, center_y + dy * scale / 2

        drawing.setStrokeColor(colors.HexColor("#7a8790"))
        drawing.setLineWidth(0.7)
        drawing.roundRect(0, 2, self.width, self.height - 4, 3, stroke=1, fill=0)
        drawing.setStrokeColor(colors.black)
        drawing.setFillColor(colors.black)
        drawing.setLineWidth(2.2)
        drawing.line(start_x, start_y, end_x, end_y)
        drawing.circle(start_x, start_y, 4, stroke=1, fill=1)
        drawing.circle(end_x, end_y, 4, stroke=1, fill=1)

        drawing.setFont("Helvetica-Bold", 8)
        drawing.drawCentredString(start_x, start_y - 16, f"N{self.node_i} (I)")
        drawing.drawCentredString(end_x, end_y - 16, f"N{self.node_j} (J)")
        drawing.setFont("Helvetica", 8)
        drawing.drawCentredString(center_x, self.height - 15, f"Barra {self.element_id} · θ = {np.degrees(self.phi):.2f}°")

        arrow_length = 18
        for node_id, x, y in ((self.node_i, start_x, start_y), (self.node_j, end_x, end_y)):
            dof_x = 2 * (node_id - 1) + 1
            dof_y = dof_x + 1
            self._arrow(x, y, x + arrow_length, y, f"g{dof_x} +X")
            self._arrow(x, y, x, y + arrow_length, f"g{dof_y} +Y")

        unit_x, unit_y = dx / max(length, 1e-9), dy / max(length, 1e-9)
        local_start_x = start_x + unit_x * length * scale * 0.18
        local_start_y = start_y + unit_y * length * scale * 0.18
        self._arrow(local_start_x, local_start_y, local_start_x + unit_x * 24, local_start_y + unit_y * 24, "+x_l")


def _draw_page(canvas, document) -> None:
    canvas.saveState()
    width, _ = landscape(A4)
    canvas.setStrokeColor(colors.HexColor("#b7c7d1"))
    canvas.line(28, 27, width - 28, 27)
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#526573"))
    canvas.drawString(28, 15, "Análisis Estructural II · Método Matricial de Rigidez")
    canvas.drawRightString(width - 28, 15, f"Página {document.page}")
    canvas.restoreState()


def export_analysis_excel(results: Dict[str, Any], output_path: Any = "Memoria_Analisis_Estructural.xlsx") -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Nodos"
    length_unit = results.get("length_unit", "m")
    force_unit = results.get("force_unit", "N")
    _records_table(ws, "NODOS", ["Nodo", f"X ({length_unit})", f"Y ({length_unit})"], [
        {"Nodo": row.get("Nodo"), f"X ({length_unit})": row.get("X"), f"Y ({length_unit})": row.get("Y")}
        for row in results.get("nodes", [])
    ])

    ws_barras = wb.create_sheet("Barras")
    _records_table(ws_barras, "PROPIEDADES DE BARRAS", ["Barra", "Nodo_i", "Nodo_j", f"E ({force_unit}/{length_unit}²)", f"A ({length_unit}²)"], [
        {
            "Barra": row.get("Barra"), "Nodo_i": row.get("Nodo_i"), "Nodo_j": row.get("Nodo_j"),
            f"E ({force_unit}/{length_unit}²)": row.get("E"), f"A ({length_unit}²)": row.get("A"),
        }
        for row in results.get("elements", [])
    ])

    ws_apoyos = wb.create_sheet("Apoyos")
    _records_table(ws_apoyos, "CONDICIONES DE APOYO", ["Nodo", "Ux", "Uy"], results.get("supports", []))

    ws_cargas = wb.create_sheet("Cargas")
    _records_table(ws_cargas, "CARGAS NODALES", ["Nodo", f"Fx ({force_unit})", f"Fy ({force_unit})"], [
        {"Nodo": item.get("Nodo"), f"Fx ({force_unit})": item.get("Fx"), f"Fy ({force_unit})": item.get("Fy")} for item in results.get("loads", [])
    ])

    ws_global = wb.create_sheet("K_Global")
    if "K_global" in results:
        dof_labels = [f"g{index + 1}" for index in range(len(results["K_global"]))]
        ws_global.append(["MATRIZ DE RIGIDEZ GLOBAL [K] (N/m)"])
        ws_global.append(["GDL", *dof_labels])
        for label, row in zip(dof_labels, results["K_global"]):
            ws_global.append([label, *[float(value) for value in row]])

    ws_vector = wb.create_sheet("Vector_F")
    if "F_global" in results:
        ws_vector.append(["CARGAS GLOBALES [F] (N)"])
        ws_vector.append(["GDL", "F (N)"])
        for index, value in enumerate(results["F_global"]):
            ws_vector.append([f"g{index + 1}", float(value)])

    if "partition" in results:
        partition_sheet = wb.create_sheet("Particion")
        free = results.get("free_dofs", [])
        restrained = results.get("restrained_dofs", [])
        labels = [f"g{index + 1}" for index in range(len(results["K_global"]))]
        for key in ("K_LL", "K_LR", "K_RL", "K_RR"):
            matrix = results["partition"][key]
            row_ids = free if key[2] == "L" else restrained
            column_ids = free if key[3] == "L" else restrained
            partition_sheet.append([f"SUBMATRIZ [{key}] (N/m)"])
            partition_sheet.append(["GDL", *[labels[index] for index in column_ids]])
            for label, row in zip((labels[index] for index in row_ids), matrix):
                partition_sheet.append([label, *[float(value) for value in row]])
            partition_sheet.append([])
        k_ll_inverse = np.linalg.inv(results["partition"]["K_LL"])
        partition_sheet.append(["INVERSA [K_LL]^-1 (m/N)"])
        partition_sheet.append(["GDL", *[labels[index] for index in free]])
        for label, row in zip((labels[index] for index in free), k_ll_inverse):
            partition_sheet.append([label, *[float(value) for value in row]])

    if "displacements" in results:
        ws_results = wb.create_sheet("Resultados")
        length_factor = results.get("length_factor", 1.0)
        force_factor = results.get("force_factor", 1.0)
        ws_results.append(["DESPLAZAMIENTOS Y REACCIONES"])
        ws_results.append(["Nodo", f"Dx ({length_unit})", f"Dy ({length_unit})", f"Rx ({force_unit})", f"Ry ({force_unit})"])
        reaction_lookup = {row["Nodo"]: row for row in results.get("reactions", [])}
        for node in results.get("nodes", []):
            node_id = int(node["Nodo"])
            dx, dy = results["displacements"][2 * (node_id - 1):2 * node_id]
            reaction = reaction_lookup.get(node_id, {"Rx": 0.0, "Ry": 0.0})
            ws_results.append([node_id, dx / length_factor, dy / length_factor, reaction["Rx"] / force_factor, reaction["Ry"] / force_factor])
        ws_results.append([])
        ws_results.append(["FUERZAS AXIALES POR BARRA"])
        ws_results.append(["Barra", "Nodos", f"N ({force_unit})", "Clasificación"])
        elements = {int(item["Barra"]): item for item in results.get("elements", [])}
        for result in results.get("axial_forces", []):
            element = elements[result["Barra"]]
            ws_results.append([result["Barra"], f"N{element['Nodo_i']}–N{element['Nodo_j']}", result["N"] / force_factor, result["estado"]])

    ws_validation = wb.create_sheet("Validacion")
    ws_validation.append(["VALIDACIÓN", "Estado"])
    ws_validation.append(["Estado", results.get("validation", {}).get("status", "PENDING")])
    ws_validation.append(["Mensaje", results.get("validation", {}).get("message", "")])

    _format_workbook(wb)
    wb.save(output_path)
    return output_path


def generate_pdf_report(results: Dict[str, Any], output_path: str = "Memoria_Analisis_Estructural.pdf") -> str:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ReportTitle", parent=styles["Title"], fontSize=18, leading=22, alignment=1, textColor=colors.HexColor("#183247"), spaceAfter=5)
    subtitle_style = ParagraphStyle("ReportSubtitle", parent=styles["Normal"], fontSize=10, leading=14, alignment=1, textColor=colors.HexColor("#526573"), spaceAfter=15)
    section_style = ParagraphStyle("ReportSection", parent=styles["Heading2"], fontSize=12, leading=15, textColor=colors.HexColor("#24445c"), spaceBefore=12, spaceAfter=7)
    subhead_style = ParagraphStyle("ReportSubhead", parent=styles["Heading3"], fontSize=9, leading=12, textColor=colors.HexColor("#35566d"), spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle("ReportBody", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#27343b"))
    available_width = landscape(A4)[0] - 56
    force_unit = results.get("force_unit", "N")
    length_unit = results.get("length_unit", "m")
    force_factor = results.get("force_factor", 1.0)
    length_factor = results.get("length_factor", 1.0)
    dof_labels = [f"g{index + 1}" for index in range(len(results.get("K_global", [])))]
    story = [
        Paragraph("ANÁLISIS ESTRUCTURAL II", title_style),
        Paragraph("Memoria de cálculo · Método matricial de rigidez", subtitle_style),
        Paragraph(str(results.get("name", "Estructura analizada")), styles["Heading1"]),
        Paragraph(f"Unidades de entrada: {length_unit} y {force_unit}. Cálculo interno realizado en SI.", body_style),
        Spacer(1, 10),
    ]

    story.append(Paragraph("1. Datos de entrada", section_style))
    for title, columns, records in (
        ("Nodos", ["Nodo", f"X ({length_unit})", f"Y ({length_unit})"], results.get("nodes", [])),
        ("Barras", ["Barra", "Nodo_i", "Nodo_j", "E", "A"], results.get("elements", [])),
        ("Apoyos", ["Nodo", "Ux", "Uy"], results.get("supports", [])),
        ("Cargas nodales", ["Nodo", "Fx", "Fy"], results.get("loads", [])),
    ):
        story.append(Paragraph(title, subhead_style))
        table_data = [columns]
        table_data.extend([[str(record.get(column, "")) for column in columns] for record in records])
        table = Table(table_data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24445c")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#aab8c2")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fa")]),
        ]))
        story.extend([table, Spacer(1, 6)])

    story.append(Paragraph("2. Desglose matricial por barra", section_style))
    element_lookup = {int(item["Barra"]): item for item in results.get("elements", [])}
    node_lookup = {int(item["Nodo"]): item for item in results.get("nodes", [])}
    axial_results = results.get("axial_forces", [])
    for result in axial_results:
        element = element_lookup[result["Barra"]]
        node_i, node_j = int(element["Nodo_i"]), int(element["Nodo_j"])
        local_labels = [f"u{node_i}'", f"v{node_i}'", f"u{node_j}'", f"v{node_j}'"]
        global_indices = [2 * (node_i - 1), 2 * (node_i - 1) + 1, 2 * (node_j - 1), 2 * (node_j - 1) + 1]
        global_labels = [dof_labels[index] for index in global_indices]
        angle = np.degrees(result["phi"])
        axial_rigidity = float(element["E"]) * float(element["A"]) / (result["Longitud"] / length_factor)
        story.append(Paragraph(
            f"Barra {result['Barra']} · N{node_i} → N{node_j} · L={result['Longitud'] / length_factor:.4f} {length_unit} · "
            f"θ={angle:.2f}° · A={float(element['A']):.5g} {length_unit}² · "
            f"E={float(element['E']):.5g} {force_unit}/{length_unit}² · AE/L={axial_rigidity:.5g} {force_unit}/{length_unit}",
            subhead_style,
        ))
        node_i_data, node_j_data = node_lookup[node_i], node_lookup[node_j]
        story.append(MemberDiagram(
            result["Barra"], node_i, node_j,
            float(node_i_data["X"]), float(node_i_data["Y"]),
            float(node_j_data["X"]), float(node_j_data["Y"]),
            result["phi"], available_width,
        ))
        story.append(Paragraph("[K_g] = [T]ᵀ · [k_l] · [T]", body_style))
        for title, matrix, rows, columns in (
            ("[T]ᵀ", result["Tg"].T, global_labels, local_labels),
            ("Matriz de rigidez local [k_l] (N/m)", result["KL"], local_labels, local_labels),
            ("[T]", result["Tg"], local_labels, global_labels),
            ("Matriz de rigidez global [K_g] (N/m)", result["Kg"], global_labels, global_labels),
        ):
            story.append(Paragraph(title, body_style))
            story.append(_matrix_table(matrix, rows, columns, available_width))
            story.append(Spacer(1, 5))

    story.append(PageBreak())
    story.append(Paragraph("3. Matriz global ensamblada [K]", section_style))
    story.append(_matrix_table(results["K_global"], dof_labels, dof_labels, available_width))

    story.append(Paragraph("4. Partición del sistema", section_style))
    partition = results.get("partition", {})
    free_ids = results.get("free_dofs", [])
    restrained_ids = results.get("restrained_dofs", [])
    for key in ("K_LL", "K_LR", "K_RL", "K_RR"):
        row_ids = free_ids if key[2] == "L" else restrained_ids
        column_ids = free_ids if key[3] == "L" else restrained_ids
        story.append(Paragraph(f"Submatriz [{key}] (N/m)", subhead_style))
        story.append(_matrix_table(partition[key], [dof_labels[index] for index in row_ids], [dof_labels[index] for index in column_ids], available_width))
        story.append(Spacer(1, 6))
    if partition.get("K_LL") is not None:
        story.append(Paragraph("Inversa de [K_LL] (m/N)", subhead_style))
        story.append(_matrix_table(np.linalg.inv(partition["K_LL"]), [dof_labels[index] for index in free_ids], [dof_labels[index] for index in free_ids], available_width))
    story.append(Paragraph("{D_L} = [K_LL]⁻¹({F_L} − [K_LR]{D_R})", body_style))
    story.append(Paragraph("{F_RR} = [K_RL]{D_L} + [K_RR]{D_R} − {F_R}", body_style))

    story.append(Paragraph("5. Desplazamientos y reacciones", section_style))
    displacement_rows = [["Nodo", f"Dx ({length_unit})", f"Dy ({length_unit})"]]
    reaction_rows = [["Nodo", f"Rx ({force_unit})", f"Ry ({force_unit})"]]
    for node in results.get("nodes", []):
        node_id = int(node["Nodo"])
        dx, dy = results["displacements"][2 * (node_id - 1):2 * node_id]
        displacement_rows.append([f"N{node_id}", f"{dx / length_factor:.6f}", f"{dy / length_factor:.6f}"])
        reaction = next((item for item in results.get("reactions", []) if item["Nodo"] == node_id), {"Rx": 0.0, "Ry": 0.0})
        reaction_rows.append([f"N{node_id}", f"{reaction['Rx'] / force_factor:.5f}", f"{reaction['Ry'] / force_factor:.5f}"])
    for table_data in (displacement_rows, reaction_rows):
        table = Table(table_data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24445c")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (1, 1), (-1, -1), "Courier"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#aab8c2")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fa")]),
        ]))
        story.extend([table, Spacer(1, 8)])

    story.append(Paragraph("6. Desarrollo de fuerzas internas por barra", section_style))
    for result in axial_results:
        element = element_lookup[result["Barra"]]
        node_i, node_j = int(element["Nodo_i"]), int(element["Nodo_j"])
        local_labels = [f"u{node_i}'", f"v{node_i}'", f"u{node_j}'", f"v{node_j}'"]
        global_indices = [2 * (node_i - 1), 2 * (node_i - 1) + 1, 2 * (node_j - 1), 2 * (node_j - 1) + 1]
        global_labels = [dof_labels[index] for index in global_indices]
        story.append(Paragraph(f"Barra {result['Barra']} · N{node_i} → N{node_j}", subhead_style))
        story.append(Paragraph("{d_l} = [T]{d_g};  {f_L} = [k_l]{d_l} = [k_l][T]{d_g}", body_style))
        recovery_data = [["GDL global", "GDL local", "d_g (m)", "d_l (m)", "f_L (N)", "Carga eq. (N)", "Fuerza interna (N)"]]
        for index, (global_label, local_label) in enumerate(zip(global_labels, local_labels)):
            recovery_data.append([
                global_label,
                local_label,
                f"{result['De'][index]:.6e}",
                f"{result['d_local'][index]:.6e}",
                f"{result['f_L'][index]:.5f}",
                f"{result['carga_equivalente_local'][index]:.5f}",
                f"{result['fuerzas_internas_locales'][index]:.5f}",
            ])
        recovery_table = Table(recovery_data, repeatRows=1, hAlign="LEFT")
        recovery_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24445c")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (2, 1), (-1, -1), "Courier"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#aab8c2")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fa")]),
        ]))
        story.extend([
            recovery_table,
            Paragraph(
                f"N = {result['N'] / force_factor:+.5f} {force_unit} · {result['estado']}",
                body_style,
            ),
            Spacer(1, 6),
        ])

    story.append(Paragraph("7. Resumen de fuerzas axiales internas", section_style))
    axial_table = [["Barra", "Nodos", f"N ({force_unit})", "Clasificación"]]
    for result in axial_results:
        element = element_lookup[result["Barra"]]
        axial_table.append([
            f"B{result['Barra']}", f"N{element['Nodo_i']}–N{element['Nodo_j']}",
            f"{result['N'] / force_factor:+.5f}", result["estado"],
        ])
    table = Table(axial_table, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24445c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (2, 1), (2, -1), "Courier"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#aab8c2")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fa")]),
    ]))
    story.extend([table, Spacer(1, 8), Paragraph("Convención: N positivo = tracción; N negativo = compresión.", body_style)])
    validation = results.get("validation", {})
    if validation:
        story.append(Paragraph("8. Validación", section_style))
        story.append(Paragraph(f"Estado: {validation.get('status', 'PENDING')} · {validation.get('message', '')}", body_style))

    doc = SimpleDocTemplate(
        output_path, pagesize=landscape(A4),
        leftMargin=28, rightMargin=28, topMargin=30, bottomMargin=38,
        title=f"Memoria de cálculo - {results.get('name', 'Estructura')}",
        author="Análisis Estructural II",
    )
    doc.build(story, onFirstPage=_draw_page, onLaterPages=_draw_page)
    return output_path


def export_excel_bytes(results: Dict[str, Any]) -> bytes:
    bio = io.BytesIO()
    export_analysis_excel(results, output_path=bio)
    bio.seek(0)
    workbook = load_workbook(bio)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    bio.close()
    bio = output
    bio.seek(0)
    return bio.read()
